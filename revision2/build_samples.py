"""Generate fictional office examples with inspectable metadata; no real records."""
from pathlib import Path
import csv,zipfile
from openpyxl import Workbook
from openpyxl.comments import Comment
from docx import Document
R=Path(__file__).resolve().parent/'downloads'
rows=[['Неделя 1','Письма',12,'Учебная группа А','a@example.invalid'],['Неделя 1','Записки',4,'Учебная группа Б','b@example.invalid'],['Неделя 2','Письма',15,'Учебная группа А','a@example.invalid'],['Неделя 2','Записки',5,'Учебная группа Б','b@example.invalid']]
headers=['Неделя','Вид работы','Количество','Условная группа','Учебный контакт']
for name,minimal in [('synthetic-data.csv',False),('synthetic-minimal.csv',True)]:
    with (R/name).open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.writer(f,delimiter=';');w.writerow(headers[:3] if minimal else headers)
        w.writerows([r[:3] if minimal else r for r in rows])
b=Workbook();s=b.active;s.title='Недельная нагрузка';s.append(['СИНТЕТИЧЕСКИЕ ДАННЫЕ — не сведения о реальной организации']);s.append(headers[:3])
for r in rows:s.append(r[:3])
s['B3'].comment=Comment('Учебное примечание: дополнительные сведения остаются в файле, даже если они не видны в ячейке.','Учебный автор')
h=b.create_sheet('Скрытый учебный лист');h.append(['Условная группа','Учебный контакт']);h.append(['Учебная группа А','a@example.invalid']);h.sheet_state='hidden';b.properties.creator='Учебный автор';b.save(R/'synthetic-file-layers.xlsx')
d=Document();d.add_heading('Отчёт по обучению',0);d.add_paragraph('СИНТЕТИЧЕСКИЙ ПРИМЕР. Не является отчётом организации.')
d.add_paragraph('Запланировано 4 занятия. Проведено 3. Одно занятие перенесено из-за изменения графика. Новая дата не согласована.')
d.add_paragraph('Для упражнения подготовьте выжимку: результат, затруднение, следующий шаг. Не придумывайте дату.')
d.core_properties.author='Учебный автор';d.core_properties.title='Отчёт по обучению — учебный пример';d.core_properties.subject='Показ свойств документа';d.core_properties.comments='Учебное описание в свойствах; это не комментарий рецензента внутри текста.'
d.save(R/'synthetic-note.docx')
with zipfile.ZipFile(R/'synthetic-note.docx') as z:
    xml=z.read('docProps/core.xml').decode('utf-8');assert 'Учебный автор' in xml
(R/'synthetic-files-readme.md').write_text('''# Учебные файлы: что именно показать
Все числа и записи созданы заново. Домены example.invalid не являются контактами сотрудников. Файлы не загружайте в онлайн-конвертеры.

## DOCX: метаданные на настоящем файле
1. Скачайте synthetic-note.docx. Это короткий учебный фрагмент, а не полный 28-страничный отчёт из сценария.
2. Откройте в настольном Word для Windows: Файл → Сведения → Свойства → Дополнительные свойства. На вкладке «Документ» (Summary) найдите автора и название.
3. В тексте нет имени, но свойство «Автор» содержит «Учебный автор». Изменение текста само по себе не очищает это поле.
4. DOCX — пакет с внутренними файлами. Для технического пояснения: поле автора записано в docProps/core.xml. Разбирать ZIP во время урока не требуется. Даты Проводника также относятся к файловой системе и могут отличаться от свойств документа.
5. Сохраните отдельную копию. Файл → Сведения → Поиск проблем → Инспектор документов → Проверить. Просмотрите категории, удалите ненужное, повторите проверку. В нашем DOCX специально заполнены свойства; комментариев рецензента и исправлений в нём нет.
6. Это не гарантия обнаружения всех скрытых сведений и не разрешение передавать рабочий файл.

## XLSX: скрытые сведения — отдельная проверка
Видимый лист «Недельная нагрузка»: четыре строки с количеством документов. Ячейка B3 содержит учебное примечание. Правой кнопкой по ярлыку листа → Показать → «Скрытый учебный лист»: появятся условные контакты.
Не удаляйте строки и листы в оригиналах без проверки зависимостей формул. Для показа используйте только копию учебного файла.

## CSV: итог и контроль
synthetic-data.csv — учебный полный набор. synthetic-minimal.csv — только неделя, вид работы, количество. Никаких минут и кабинетов.
Неделя 1: 12 + 4 = 16. Неделя 2: 15 + 5 = 20. Письма: 27; записки: 9; всего 36.
Эти числа показывают объём документов, но не качество и не эффективность сотрудников.

## Источники команд интерфейса
Microsoft: https://support.microsoft.com/en-us/office/collab-files/view-or-change-the-properties-for-an-office-file
Microsoft: https://support.microsoft.com/en-us/office/collab-files/remove-hidden-data-and-personal-information-by-inspecting-documents-presentations-or-workbooks
Пути относятся к настольному Office для Windows; названия могут отличаться по версии и локализации. В других редакторах найдите соответствующие команды по их справке.
''',encoding='utf-8')
print('Created DOCX with verified author metadata; XLSX with hidden sheet; 4-row workload CSVs')
