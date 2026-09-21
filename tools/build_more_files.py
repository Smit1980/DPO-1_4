# -*- coding: utf-8 -*-
"""Дополнительные файлы урока 4: пакет промптов (docx/xlsx/csv/json), бланк ДЗ, журнал заявок 50 строк, слайды PDF и PPTX."""
from __future__ import annotations

import csv
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import prompts as P  # noqa: E402
from lib import SLIDES, img_name  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "data" / "urok4"
DISCLAIMER = ("Материал не заменяет политику информационной безопасности Концерна ВКО «Алмаз – Антей» и локальные акты предприятия. "
              "Все данные в промптах условные.")

# Порядок показа на занятии
TEACH_ORDER = [0, 21] + list(range(1, 11)) + [22] + list(range(11, 21)) + [23]

GROUP_FILL = {"Подготовка данных": "E5EFFF", "Шаблоны документов": "DCFCE7", "Роль в промпте": "FEF3C7", "Проверка результата": "FDE2E2", "Практикум 1.3": "EDE9FE"}


def prompt_pack() -> None:
    import docx
    from docx.shared import Pt, RGBColor
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    allp = P.PROMPTS + P.PRACTICE_PROMPTS
    (OUT / "prompty_urok4.txt").write_text(P.prompt_txt(), encoding="utf-8")
    with (OUT / "prompty_urok4.csv").open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["id", "group", "title", "description", "prompt"])
        for p in allp:
            w.writerow([p["id"], p["group"], p["title"], p.get("description", ""), p["content"]])
    (OUT / "prompty_urok4.json").write_text(json.dumps([{k: p.get(k, "") for k in ("id", "group", "title", "description", "content")} for p in allp], ensure_ascii=False, indent=1), encoding="utf-8")

    d = docx.Document()
    d.core_properties.author = "ДПО-1 · Клуб «ИИ в образовании»"
    d.core_properties.title = "Промпты урока 4"
    d.add_heading("Промпты урока 4 · ДПО-1", 0)
    d.add_paragraph(DISCLAIMER)
    d.add_paragraph("Всё, что зависит от фактов, стоит в квадратных скобках — заполняете вы. Роль в промпте задаёт формат, тон и позицию; проверка фактов остаётся за вами.")
    group = None
    for p in allp:
        if p["group"] != group:
            group = p["group"]
            d.add_heading(group, 1)
        d.add_heading(f'{p["id"]} · {p["title"]}', 2)
        if p.get("description"):
            para = d.add_paragraph(p["description"])
            para.runs[0].italic = True
        for line in p["content"].split("\n"):
            para = d.add_paragraph(line)
            para.paragraph_format.space_after = Pt(2)
            for r in para.runs:
                r.font.name = "Consolas"
                r.font.size = Pt(10)
    d.save(OUT / "prompty_urok4.docx")

    wb = Workbook()
    ws = wb.active
    ws.title = "Промпты"
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:E1")
    ws["A1"] = "Промпты урока 4 · ДПО-1"
    ws["A1"].font = Font(name="Calibri", size=18, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="002060")
    ws["A1"].alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 34
    ws.merge_cells("A2:E2")
    ws["A2"] = DISCLAIMER
    ws["A2"].font = Font(italic=True, size=10, color="002060")
    ws["A2"].fill = PatternFill("solid", fgColor="E5EFFF")
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="center", indent=1)
    ws.row_dimensions[2].height = 30
    line = Side(style="thin", color="C7D2E5")
    border = Border(left=line, right=line, top=line, bottom=line)
    heads = ["№", "Группа", "Название", "Описание", "Промпт (копируйте ячейку целиком)"]
    widths = [7, 22, 36, 42, 110]
    for j, (h, wd) in enumerate(zip(heads, widths), 1):
        c = ws.cell(4, j, h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0047FF")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border
        ws.column_dimensions[chr(64 + j)].width = wd
    for i, p in enumerate(allp, 5):
        vals = [p["id"], p["group"], p["title"], p.get("description", ""), p["content"]]
        for j, v in enumerate(vals, 1):
            c = ws.cell(i, j, v)
            c.alignment = Alignment(wrap_text=True, vertical="top", horizontal="center" if j == 1 else "left")
            c.border = border
            c.font = Font(name="Consolas" if j == 5 else "Calibri", size=10)
            if j <= 3:
                c.fill = PatternFill("solid", fgColor=GROUP_FILL.get(p["group"], "FFFFFF"))
        ws.row_dimensions[i].height = min(15 * (p["content"].count("\n") + 2 + len(p["content"]) // 120), 260)
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:E{4 + len(allp)}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(OUT / "prompty_urok4.xlsx")
    print(f"  ✓ пакет промптов: {len(P.PROMPTS)} + {len(P.PRACTICE_PROMPTS)} практикума (docx, xlsx, csv, json, txt)")


def hw_blank() -> None:
    import docx
    d = docx.Document()
    d.core_properties.author = "ДПО-1 · Клуб «ИИ в образовании»"
    d.core_properties.title = "Домашнее задание · урок 4"
    d.add_heading("Домашнее задание · урок 4", 0)
    d.add_paragraph("Реальных названий, фамилий и рабочих данных не записывайте. " + DISCLAIMER)
    for n in (1, 2):
        d.add_heading(f"Часть 1.{n}. Паспорт безопасной задачи", 1)
        t = d.add_table(rows=0, cols=2)
        t.style = "Table Grid"
        for a in ["Тип задачи (одной фразой)", "Что нужно ИИ", "Что убрать", "Что заменить или обобщить", "Где обрабатывать (разрешённый инструмент)", "Кто проверит результат", "Что я не знаю и у кого уточню"]:
            row = t.add_row().cells
            row[0].text, row[1].text = a, ""
    d.add_heading("Часть 2. Пять проверок по одному результату ИИ", 1)
    t = d.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    for c, h in zip(t.rows[0].cells, ["Проверка", "Что проверил", "Что нашёл и что исключил"]):
        c.text = h
    for a in ["1. Достоверность", "2. Источник", "3. Права и согласие", "4. Предвзятость и последствия", "5. Ответственность"]:
        r = t.add_row().cells
        r[0].text = a
    d.add_heading("Часть 3. Вопросы ответственному по ИБ или юристам", 1)
    for i in range(1, 9):
        d.add_paragraph(f"{i}. ____________________________________________")
    d.add_heading("Часть 4. Фраза-завершение", 1)
    d.add_paragraph("Перед загрузкой информации в ИИ я сначала буду… ____________________")
    d.add_paragraph("Перед использованием результата ИИ я сначала буду… ____________________")
    d.save(OUT / "domashnee_zadanie_blank.docx")
    print("  ✓ бланк домашнего задания")


def journal50() -> None:
    """Учебный журнал на 50 строк. Свойства файла очищаются от реальных имён; ловушка (инструкция для ИИ в ячейке J16) сохраняется намеренно."""
    import datetime as dt
    import openpyxl
    src = ROOT / "data" / "sinteticheskie_zayavki_50.xlsx"
    if not src.is_file():
        print("  ! нет исходного журнала на 50 строк")
        return
    wb = openpyxl.load_workbook(src)
    wb.properties.creator = "Комарова Е.В. (учебный)"
    wb.properties.lastModifiedBy = "Ершов Н.К. (учебный)"
    wb.properties.title = "Журнал заявок АХО · 50 строк (учебный)"
    wb.properties.created = dt.datetime(2026, 9, 2, 9, 0)
    wb.properties.modified = dt.datetime(2026, 9, 22, 17, 0)
    wb.save(OUT / "zhurnal_zayavok_50_uchebnyj.xlsx")
    print("  ✓ журнал заявок на 50 строк (свойства очищены)")


def slides_pdf_pptx() -> None:
    from PIL import Image
    src = ROOT / "png" / "urok4"
    imgs = []
    for n in TEACH_ORDER:
        f = src / ("00.png" if n == 0 else f"{n}.png")
        if f.is_file():
            im = Image.open(f).convert("RGB")
            im = im.resize((1600, round(im.height * 1600 / im.width)))
            imgs.append((n, im))
    if not imgs:
        print("  ! слайдов нет")
        return
    imgs[0][1].save(OUT / "slaydy_urok4.pdf", save_all=True, append_images=[i for _, i in imgs[1:]], resolution=110.0, quality=82)
    try:
        from pptx import Presentation
        from pptx.util import Emu
        import io
        prs = Presentation()
        prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
        blank = prs.slide_layouts[6]
        for n, im in imgs:
            sl = prs.slides.add_slide(blank)
            buf = io.BytesIO()
            im.save(buf, "JPEG", quality=88)
            buf.seek(0)
            sl.shapes.add_picture(buf, 0, 0, width=prs.slide_width, height=prs.slide_height)
            sl.notes_slide.notes_text_frame.text = f"Слайд {n:02d}. {SLIDES[n][1]}. Пояснение — на странице урока и в конспекте."
        prs.core_properties.title = "Урок 4 · Подготовка информации для ИИ и правовые риски"
        prs.core_properties.author = "ДПО-1 · Клуб «ИИ в образовании»"
        prs.save(OUT / "slaydy_urok4.pptx")
        print(f"  ✓ презентация: {len(imgs)} слайдов (pdf, pptx)")
    except ImportError:
        print(f"  ✓ pdf: {len(imgs)} слайдов; pptx пропущен (нет python-pptx)")


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    prompt_pack()
    hw_blank()
    journal50()
    slides_pdf_pptx()


if __name__ == "__main__":
    main()
