# -*- coding: utf-8 -*-
"""Учебные файлы урока 4 → data/urok4/ и архив data/urok4_materialy.zip. Запуск: py tools/build_files.py
Все данные вымышлены. Нужны: python-docx, openpyxl, Pillow; для PDF — Chrome или Edge."""
from __future__ import annotations

import html
import shutil
import subprocess
import sys
import tempfile
import time
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import prompts as P  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "data" / "urok4"
E = html.escape

DISCLAIMER = ("Памятка не заменяет политику информационной безопасности Концерна ВКО «Алмаз – Антей» и локальные акты предприятия. "
              "Все примеры вымышлены.")


# ── 1. DOCX с настоящими скрытыми данными ───────────────────────────────────
def build_metadata_docx() -> None:
    import datetime as dt
    import docx
    from docx.shared import Pt

    d = docx.Document()
    cp = d.core_properties
    cp.author = "Комарова Е.В. (учебный)"
    cp.last_modified_by = "Ершов Н.К. (учебный)"
    cp.title = "Договор_45-2026_ЧЕРНОВИК"
    cp.subject = "Перенос срока поставки"
    cp.comments = "Внутренний номер проекта ВН-2026/118. Не отправлять поставщику."
    cp.created = dt.datetime(2026, 9, 2, 9, 15)
    cp.modified = dt.datetime(2026, 9, 3, 17, 40)
    sec = d.sections[0]
    sec.header.paragraphs[0].text = "Для служебного пользования (учебный пример) · Проект ВН-2026/118"
    d.add_heading("Уведомление о переносе срока поставки", level=1)
    p = d.add_paragraph("Уважаемые партнёры! Сообщаем, что срок поставки по договору сдвигается на две недели. Просим подтвердить новую дату. ")
    r = p.add_run("Скидку 7 % даём только при предоплате.")
    r.font.hidden = True
    d.add_paragraph("С уважением, руководитель отдела закупок.")
    d.add_paragraph("Учебный файл к практике 2 урока 4. Все данные вымышлены. Найдите скрытое: свойства, примечание, скрытый текст, колонтитул.").runs[0].font.size = Pt(9)
    try:
        d.add_comment(runs=p.runs[0], text="Цену не показывать, согласовать с Ершовым", author="Ершов Н.К. (учебный)", initials="ЕН")
    except Exception as exc:  # старые версии python-docx
        print("  ! примечание не добавлено:", exc)
    OUT.mkdir(parents=True, exist_ok=True)
    path = OUT / "zayavka_uchebnaya_s_metadannymi.docx"
    d.save(path)
    # Организация — в app.xml (свойство «Организация» в Word)
    tmp = path.with_suffix(".tmp")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "docProps/app.xml":
                s = data.decode("utf-8")
                if "<Company>" in s:
                    import re
                    s = re.sub(r"<Company>.*?</Company>", "<Company>Учебное предприятие</Company>", s)
                else:
                    s = s.replace("</Properties>", "<Company>Учебное предприятие</Company></Properties>")
                data = s.encode("utf-8")
            zout.writestr(item, data)
    tmp.replace(path)
    print("  ✓", path.name)


# ── 2. XLSX: журнал заявок со скрытым листом ────────────────────────────────
JOURNAL = [
    (1, "Комарова Е.В.", "Отдел закупок", "+7 900 000-01-11", "Канцтовары", "02.09.2026 09:15", 214, "Папки-скоросшиватели, 20 шт."),
    (2, "Лебедев О.И.", "Планово-экономический отдел", "+7 900 000-01-12", "Пропуск гостя", "02.09.2026 11:20", 305, "Гость из организации-партнёра на 04.09"),
    (3, "Фомина Т.С.", "Главный метролог (1 чел.)", "доб. 4417", "Ремонт кабинета", "03.09.2026 08:30", 118, "Течёт потолок. Звонить только мне"),
    (4, "Ершов Н.К.", "Отдел закупок", "+7 900 000-01-13", "Заказ транспорта", "03.09.2026 17:45", "—", "Поездка 10.09 на завод в другом городе"),
    (5, "Комарова Е.В.", "Отдел закупок", "+7 900 000-01-11", "Командировка", "04.09.2026 12:00", 214, "Билеты на 15.09; почта komarova@example.test"),
    (6, "Гордеев П.А.", "Служба качества", "+7 900 000-01-14", "Канцтовары", "07.09.2026 10:05", 402, "Маркеры, 10 шт."),
    (7, "Лебедев О.И.", "Планово-экономический отдел", "+7 900 000-01-12", "Пропуск гостя", "08.09.2026 14:30", 305, "Гость на 09.09, без сопровождающего"),
    (8, "Фомина Т.С.", "Главный метролог (1 чел.)", "доб. 4417", "Канцтовары", "09.09.2026 09:40", 118, "Бумага А4, 5 пачек"),
]
SAFE = [("З-01", "Канцтовары", 1), ("З-02", "Пропуск гостя", 1), ("З-03", "Ремонт кабинета", 1), ("З-04", "Заказ транспорта", 1),
        ("З-05", "Командировка", 1), ("З-06", "Канцтовары", 2), ("З-07", "Пропуск гостя", 2), ("З-08", "Канцтовары", 2)]


def build_journal() -> None:
    import csv
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Журнал"
    ws.append(["№", "Заявитель", "Подразделение", "Телефон", "Тип заявки", "Дата и время", "Каб.", "Комментарий"])
    for row in JOURNAL:
        ws.append(list(row))
    hid = wb.create_sheet("Служебный")
    hid.append(["Учебная заметка", "Скрытый лист: внутренний номер проекта ВН-2026/118, список согласующих (вымышленный)"])
    hid.sheet_state = "hidden"
    ws.column_dimensions["G"].hidden = True
    wb.properties.creator = "Комарова Е.В. (учебный)"
    wb.properties.lastModifiedBy = "Ершов Н.К. (учебный)"
    wb.properties.title = "Журнал заявок АХО (учебный)"
    wb.save(OUT / "zhurnal_zayavok_uchebnyj.xlsx")
    with (OUT / "zhurnal_zayavok_bezopasnyj.csv").open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["Метка", "Тип заявки", "Неделя"])
        w.writerows(SAFE)
    print("  ✓ журнал заявок xlsx + csv")


# ── 3. JPG с EXIF ───────────────────────────────────────────────────────────
def build_photo() -> None:
    from PIL import Image, ImageDraw, ImageFont
    img = Image.new("RGB", (900, 600), "#e8f0ff")
    dr = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype("arial.ttf", 34)
    except Exception:
        font = ImageFont.load_default()
    dr.rectangle((60, 80, 840, 520), outline="#4338ca", width=4)
    dr.text((90, 230), "Учебное фото стенда", fill="#1e293b", font=font)
    dr.text((90, 290), "(вымышленный объект)", fill="#64748b", font=font)
    exif = Image.Exif()
    exif[0x010F] = "Учебный телефон"
    exif[0x0110] = "Модель-X"
    exif[0x0131] = "Учебная камера 1.0"
    exif[0x013B] = "Комарова Е.В. (учебный)"
    exif[0x9003] = "2026:09:02 09:15:00"
    try:
        gps = exif.get_ifd(0x8825)
        gps[1] = "N"; gps[2] = (10.0, 0.0, 0.0); gps[3] = "E"; gps[4] = (20.0, 0.0, 0.0)
        exif[0x8825] = gps
    except Exception as exc:
        print("  ! GPS не добавлен:", exc)
    img.save(OUT / "foto_uchebnoe_s_exif.jpg", quality=88, exif=exif)
    print("  ✓ foto_uchebnoe_s_exif.jpg")


# ── 4. Промпты ──────────────────────────────────────────────────────────────
def build_prompts() -> None:
    import docx
    from docx.shared import Pt
    (OUT / "prompty_urok4.txt").write_text(P.prompt_txt(), encoding="utf-8")
    d = docx.Document()
    d.core_properties.author = "ДПО-1 · Клуб «ИИ в образовании»"
    d.core_properties.title = "Промпты урока 4"
    d.add_heading("Промпты урока 4 · ДПО-1", level=1)
    d.add_paragraph("Все данные в промптах условные. Реальные документы, фамилии, суммы и реквизиты в промпты не вставляйте. " + DISCLAIMER)
    for p in P.PROMPTS:
        d.add_heading(f'{p["id"]} · {p["title"]}', level=2)
        d.add_paragraph(p["description"]).italic = True
        for line in p["content"].split("\n"):
            para = d.add_paragraph(line)
            for r in para.runs:
                r.font.name = "Consolas"
                r.font.size = Pt(10)
    d.save(OUT / "prompty_urok4.docx")
    print("  ✓ промпты txt + docx")


# ── 5. Паспорт задачи ────────────────────────────────────────────────────────
def build_passport() -> None:
    import docx
    d = docx.Document()
    d.core_properties.author = "ДПО-1 · Клуб «ИИ в образовании»"
    d.core_properties.title = "Паспорт безопасной задачи"
    d.add_heading("Паспорт безопасной задачи", level=1)
    d.add_paragraph("Заполните за две минуты перед тем, как поручить задачу ИИ. Реальные названия и данные не записывайте. " + DISCLAIMER)
    rows = [("Задача", "Одно предложение: что нужно получить", "Составить письмо поставщику о переносе срока поставки"),
            ("Что нужно ИИ", "Минимум для решения", "Тип письма, тон, структура; «поставщик А», «срок сдвинулся на две недели»"),
            ("Что убрать", "Всё, что не влияет на результат", "Телефон и почта представителя, номер договора, фамилия подписанта"),
            ("Что заменить или обобщить", "Что оставить, но в безопасном виде", "Название → «поставщик А»; сумма → «около 13 млн»; дата → «на две недели»"),
            ("Где обрабатывать", "Разрешён ли инструмент", "Только инструмент, допустимый по политике безопасности. Нет подтверждения — не отправлять"),
            ("Кто проверит результат", "Человек, а не «ещё один запрос»", "Автор письма и начальник отдела перед отправкой"),
            ("Что я не знаю и у кого уточню", "Вопрос — и функция, у которой спросить", "Можно ли использовать внешний сервис для писем контрагентам → ответственный за ИБ")]
    t = d.add_table(rows=1, cols=4)
    t.style = "Table Grid"
    for c, h in zip(t.rows[0].cells, ["Поле", "Что писать", "Образец", "Моя задача"]):
        c.text = h
    for a, b, c_, in rows:
        cells = t.add_row().cells
        cells[0].text, cells[1].text, cells[2].text = a, b, c_
    d.add_paragraph()
    d.add_paragraph("Формула: ЗАДАЧА → НЕОБХОДИМЫЕ ДАННЫЕ → МИНИМИЗАЦИЯ → ОБЕЗЛИЧИВАНИЕ / ОБОБЩЕНИЕ → ПРОВЕРКА НАСТРОЕК → ПЕРЕДАЧА В ИИ.")
    d.save(OUT / "pasport_bezopasnoy_zadachi.docx")
    print("  ✓ паспорт задачи")


# ── 6. Памятки PDF + DOCX ────────────────────────────────────────────────────
MEMOS = [
    dict(file="pamyatka_1_pered_otpravkoy_v_ii", title="Перед тем как отправить в ИИ",
         blocks=[
             ("Пять вопросов перед загрузкой", ["Зачем ИИ этот файл? Что должно получиться?", "Какой минимум нужен: фрагмент или весь документ?",
                                                "Что в нём лишнее: ФИО, телефоны, номера, комментарии, чужие данные?", "Что можно заменить или обобщить?",
                                                "Разрешён ли этот инструмент для этих сведений (политика ИБ предприятия)?"]),
             ("Четыре приёма", ["Удалить — то, что не влияет на результат.", "Заменить — имя → роль или метка («поставщик А»). Таблицу соответствия хранить отдельно.",
                                "Обобщить — точное → примерное («около 13 млн», «на две недели»).", "Условные данные — придумать новые, а не «перекрасить» реальные."]),
             ("Проверьте файл: шесть мест", ["Имя файла", "Свойства (автор, «последним сохранил», организация)", "Примечания на полях", "Исправления и история",
                                             "Скрытый текст, скрытые листы, строки и столбцы", "Колонтитулы, пути и ссылки"]),
             ("Как убрать", ["Лучше вставить в запрос текст фрагмента (через «Блокнот»), чем отправлять файл.", "Если файл нужен и это разрешено: работать с копией; Файл → Сведения → Поиск проблем → Инспектор документов.",
                             "Не просить ИИ «обезличить» оригинал: передача уже состоится."]),
             ("Если данные уже отправлены", ["Остановиться, не отправлять больше.", "Зафиксировать: сервис, время, что ушло (скриншот, не пересылка файла).",
                                             "Сразу сообщить по установленному каналу (ИБ или руководитель).", "Не рассылать копии; удаление чата не гарантирует удаления данных.", "Действовать по указаниям ответственных."]),
             ("Помните", ["«Обучение отключено» ≠ «данные удалены».", "Российский сервис ≠ автоматически разрешён.", "Роль в промпте («ты — эксперт») не даёт права на данные."]),
         ]),
    dict(file="pamyatka_2_proverka_rezultata_ii", title="Проверка результата ИИ",
         blocks=[
             ("Пять проверок", ["Достоверность — это правда?", "Источник — откуда это взялось и подтверждает ли источник именно это?", "Права и согласие — можем ли мы это использовать?",
                                "Предвзятость и последствия — кого и как может затронуть результат?", "Ответственность — кто принимает окончательное решение?"]),
             ("Как проверить факт: четыре шага", ["Найти оригинал (закон, официальный сайт, первоисточник).", "Сверить цифру, дату, название.", "Убедиться, что источник подтверждает именно это утверждение.",
                                                  "Проверить, что сам источник существует. Повторный вопрос той же модели — не проверка."]),
             ("Тест «поменяй одну деталь»", ["Два одинаковых запроса в новых чатах; меняется одна характеристика (возраст, пол, имя, город).", "Сравнить: оценку, тон, «ярлыки», рекомендацию.",
                                              "Если ответ изменился при равных деловых качествах — сигнал. Задать критерии явно и повторить.", "Решение о человеке принимает человек."]),
             ("Необычное поручение или «видео руководителя»", ["Признаки: срочность, секретность, новый канал или адрес, обход процедуры, ссылка или пароль.", "Остановиться. Подтвердить по ранее известному контакту (не по номеру из сообщения).",
                                                                "Спросить основание передачи. Сообщить в ИБ, если подозрение остаётся.", "Голос и лицо можно подделать: «похоже» — не доказательство."]),
             ("Кто отвечает", ["Тот, кто подписал и отправил документ. «Так написала нейросеть» — не аргумент.", "Сообщать об ошибке нужно сразу."]),
         ]),
]


def find_browser() -> str | None:
    for c in (r"C:\Program Files\Google\Chrome\Application\chrome.exe", r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
              r"C:\Program Files\Microsoft\Edge\Application\msedge.exe"):
        if Path(c).is_file():
            return c
    return shutil.which("chrome") or shutil.which("msedge")


def memo_html(m: dict) -> str:
    body = "".join(f"<h2>{E(t)}</h2><ul>" + "".join(f"<li>{E(i)}</li>" for i in items) + "</ul>" for t, items in m["blocks"])
    return (f'<!DOCTYPE html><html lang="ru"><head><meta charset="utf-8"><title>{E(m["title"])}</title><style>'
            '@page{size:A4;margin:14mm}body{font-family:"Segoe UI",Arial,sans-serif;color:#0f172a;font-size:11.5pt;line-height:1.4}'
            'h1{font-size:20pt;color:#4338ca;margin:0 0 4px}.sub{color:#64748b;font-size:9.5pt;margin-bottom:10px}h2{font-size:12.5pt;margin:12px 0 4px;color:#4338ca}'
            'ul{margin:0 0 4px;padding-left:20px}li{margin-bottom:2px}.foot{margin-top:14px;padding-top:8px;border-top:1px solid #cbd5e1;font-size:8.5pt;color:#64748b}'
            f'</style></head><body><h1>{E(m["title"])}</h1><div class="sub">ДПО-1 · Урок 4 · Клуб «ИИ в образовании»</div>{body}<div class="foot">{E(DISCLAIMER)}</div></body></html>')


def print_pdf(browser: str, source_html: str, pdf: Path) -> bool:
    with tempfile.TemporaryDirectory(prefix="memo_pdf_") as tmp:
        tp = Path(tmp)
        (tp / "memo.html").write_text(source_html, encoding="utf-8")
        out = tp / "memo.pdf"
        subprocess.run([browser, "--headless", "--disable-gpu", "--no-pdf-header-footer", f"--user-data-dir={tp / 'profile'}",
                        f"--print-to-pdf={out}", (tp / "memo.html").as_uri()], check=False, capture_output=True, timeout=180)
        for _ in range(30):
            if out.is_file() and out.stat().st_size > 1000:
                shutil.copyfile(out, pdf)
                return True
            time.sleep(0.5)
    return False


def build_memos() -> None:
    import docx
    browser = find_browser()
    for m in MEMOS:
        d = docx.Document()
        d.core_properties.author = "ДПО-1 · Клуб «ИИ в образовании»"
        d.core_properties.title = m["title"]
        d.add_heading(m["title"], level=1)
        for t, items in m["blocks"]:
            d.add_heading(t, level=2)
            for i in items:
                d.add_paragraph(i, style="List Bullet")
        d.add_paragraph(DISCLAIMER)
        d.save(OUT / f'{m["file"]}.docx')
        ok = bool(browser) and print_pdf(browser, memo_html(m), OUT / f'{m["file"]}.pdf')
        print("  ✓" if ok else "  ! PDF НЕ создан:", m["file"])
        if not ok:
            raise SystemExit("PDF памятки не создан — проверьте Chrome/Edge")


def build_zip() -> None:
    z = ROOT / "data" / "urok4_materialy.zip"
    with zipfile.ZipFile(z, "w", zipfile.ZIP_DEFLATED) as zf:
        for f in sorted(OUT.iterdir()):
            if f.name.startswith("slaydy_"):
                continue  # слайды — отдельными файлами (большой размер)
            zf.write(f, f.name)
    print("  ✓", z.name, len(list(OUT.iterdir())), "файлов")


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    build_metadata_docx()
    build_journal()
    build_photo()
    build_passport()
    build_memos()
    import build_more_files
    import build_xlsx
    build_more_files.main()
    print("  ✓", build_xlsx.build().name)
    build_zip()


if __name__ == "__main__":
    main()
